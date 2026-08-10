"""Report export to PDF + DOCX with classification header/footer, charts, and a
sources list, so an exported file stands on its own away from the app
(§5.10, FR-5.4.4)."""
from __future__ import annotations

import io

from dkip.core.config import settings


def _sources(report: dict) -> list[str]:
    """Unique 'DOC §section p.n' lines across every section, in first-seen order."""
    seen, out = set(), []
    for section in report.get("sections", []):
        for c in section.get("citations", []) or []:
            line = f"{c.get('doc', '')} §{c.get('section') or '-'} p.{c.get('page', '')}"
            if line not in seen:
                seen.add(line)
                out.append(line)
    return out


def _chart_table(chart: dict) -> tuple[list[str], list[list[str]]]:
    """Chart data as a printable table — exports never depend on a rendered image."""
    if chart.get("type") == "table":
        return chart.get("columns", []), [[str(c) for c in row]
                                          for row in chart.get("rows", [])]
    series = chart.get("series") or ["value"]
    header = ["", *[s.title() for s in series]]
    rows = [[str(p.get("label", "")), *[str(p.get(s, "")) for s in series]]
            for p in chart.get("points", [])]
    return header, rows


def to_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, PageTemplate,
                                    Paragraph, Spacer, Table, TableStyle)

    buf = io.BytesIO()
    banner = settings.CLASSIFICATION_BANNER
    styles = getSampleStyleSheet()

    def _chrome(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8)
        w, h = A4
        for y in (h - 12 * mm, 8 * mm):
            canvas.setFillColorRGB(0.10, 0.32, 0.18)
            canvas.rect(0, y, w, 8 * mm, fill=1, stroke=0)
            canvas.setFillColorRGB(1, 1, 1)
            canvas.drawCentredString(w / 2, y + 2.4 * mm, banner)
        canvas.restoreState()

    doc = BaseDocTemplate(buf, pagesize=A4, topMargin=22 * mm, bottomMargin=20 * mm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=_chrome)])

    flow = [Paragraph(report.get("title", "Report"), styles["Title"]), Spacer(1, 6 * mm)]
    for section in report.get("sections", []):
        flow.append(Paragraph(section["label"], styles["Heading2"]))
        flow.append(Paragraph(section.get("value", "").replace("\n", "<br/>"),
                              styles["BodyText"]))
        flow.append(Spacer(1, 4 * mm))

    images = {i["key"]: i["png"] for i in report.get("chart_images", []) if i.get("png")}
    for chart in report.get("charts", []):
        flow.append(Paragraph(chart.get("title", "Chart"), styles["Heading2"]))
        png = images.get(chart.get("title")) or (next(iter(images.values())) if images else None)
        if png and chart.get("type") != "table":
            flow.append(Image(io.BytesIO(png), width=doc.width, height=doc.width * 0.5,
                              kind="proportional"))
        header, rows = _chart_table(chart)
        if rows:
            table = Table([header, *rows], hAlign="LEFT")
            table.setStyle(TableStyle([
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.grey),
                ("LINEBELOW", (0, 0), (-1, 0), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            flow.append(Spacer(1, 3 * mm))
            flow.append(table)
        if chart.get("source") == "documents":
            flow.append(Spacer(1, 2 * mm))
            flow.append(Paragraph("Figures taken from the cited document text.",
                                  styles["Italic"]))
        flow.append(Spacer(1, 4 * mm))

    sources = _sources(report)
    if sources:
        flow.append(Paragraph("Sources", styles["Heading2"]))
        for i, line in enumerate(sources, start=1):
            flow.append(Paragraph(f"{i}. {line}", styles["BodyText"]))

    doc.build(flow)
    return buf.getvalue()


def to_docx(report: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    d = Document()
    banner = d.sections[0].header.paragraphs[0]
    banner.text = settings.CLASSIFICATION_BANNER
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = d.sections[0].footer.paragraphs[0]
    footer.text = settings.CLASSIFICATION_BANNER
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading(report.get("title", "Report"), 0)
    for section in report.get("sections", []):
        d.add_heading(section["label"], level=2)
        d.add_paragraph(section.get("value", ""))

    images = {i["key"]: i["png"] for i in report.get("chart_images", []) if i.get("png")}
    for chart in report.get("charts", []):
        d.add_heading(chart.get("title", "Chart"), level=2)
        png = images.get(chart.get("title")) or (next(iter(images.values())) if images else None)
        if png and chart.get("type") != "table":
            d.add_picture(io.BytesIO(png), width=Inches(6))
        header, rows = _chart_table(chart)
        if rows:
            table = d.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, label in zip(table.rows[0].cells, header):
                cell.text = str(label)
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
        if chart.get("source") == "documents":
            d.add_paragraph("Figures taken from the cited document text.").italic = True

    sources = _sources(report)
    if sources:
        d.add_heading("Sources", level=2)
        for i, line in enumerate(sources, start=1):
            d.add_paragraph(f"{i}. {line}")

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
